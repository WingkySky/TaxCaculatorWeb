#!/usr/bin/env python3

import sys

try:
    print('尝试导入 openpyxl...')
    import openpyxl
    wb = openpyxl.load_workbook('发放明细.xls', read_only=True, data_only=True)
    print('Sheet names:', wb.sheetnames)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f'\n=== {sheet_name} ===')
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 20:
                break
            print(row)
    sys.exit(0)
except Exception as e:
    print(f'openpyxl 失败: {e}', file=sys.stderr)
    import traceback
    traceback.print_exc()

try:
    print('\n尝试导入 xlrd...')
    import xlrd
    wb = xlrd.open_workbook('发放明细.xls')
    print('Sheet names:', wb.sheet_names())
    for sheet_idx in range(wb.nsheets):
        ws = wb.sheet_by_index(sheet_idx)
        print(f'\n=== {wb.sheet_names()[sheet_idx]} ===')
        for row_idx in range(min(ws.nrows, 20)):
            print(ws.row_values(row_idx))
    sys.exit(0)
except Exception as e:
    print(f'xlrd 失败: {e}', file=sys.stderr)
    import traceback
    traceback.print_exc()

try:
    print('\n尝试导入 pandas...')
    import pandas as pd
    df = pd.read_excel('发放明细.xls')
    print(df.to_string())
    print('\n列名:', df.columns.tolist())
    sys.exit(0)
except Exception as e:
    print(f'pandas 失败: {e}', file=sys.stderr)
    import traceback
    traceback.print_exc()

print('\n没有可用的 Excel 库！')
